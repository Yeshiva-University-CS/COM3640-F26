"""
pytest suite for pipeline_runtime.py.

Run with: uv run pytest test_pipeline_runtime.py -v
(see the accompanying setup instructions for the uv commands to create the
venv and install dependencies first)

Excel tests require DuckDB's `excel` extension, which needs a one-time
download from extensions.duckdb.org the first time it's loaded. If that
download isn't possible in your environment (offline, restricted network),
every excel-dependent test is skipped automatically rather than failing.
"""

import os
import re

import pytest

import pipeline_runtime as rt


# ---------------------------------------------------------------------------
# Fixtures: build test data files on the fly so the suite is self-contained
# ---------------------------------------------------------------------------

@pytest.fixture
def crm_csv(tmp_path):
    p = tmp_path / "crm.csv"
    p.write_text(
        "id,Full Name,status,segment\n"
        "1,Alice Smith,active,enterprise\n"
        "2,Bob Jones,inactive,smb\n"
        "3,Carol Lee,active,enterprise\n"
    )
    return str(p)


@pytest.fixture
def headerless_csv(tmp_path):
    p = tmp_path / "headerless.csv"
    p.write_text(
        "1,2.5,active\n"
        "2,3.5,inactive\n"
        "3,4.5,active\n"
    )
    return str(p)


@pytest.fixture
def billing_json(tmp_path):
    p = tmp_path / "billing.json"
    p.write_text(
        '[{"cust_id": 1, "amt_due": 50.0}, '
        '{"cust_id": 3, "amt_due": 25.0}, '
        '{"cust_id": 4, "amt_due": 10.0}]'
    )
    return str(p)


@pytest.fixture
def time_and_datetime_json(tmp_path):
    p = tmp_path / "time_and_datetime.json"
    p.write_text(
        '[{"id": 1, "event_time": "14:30:00", "event_ts": "2024-01-01T14:30:00"}, '
        '{"id": 2, "event_time": "09:00:00", "event_ts": "2024-01-02T09:00:00"}]'
    )
    return str(p)


@pytest.fixture(scope="session")
def excel_available():
    try:
        rt._ensure_excel_extension()
        return True
    except rt.PipelineRuntimeError:
        return False


@pytest.fixture
def workbook_xlsx(tmp_path, excel_available):
    if not excel_available:
        pytest.skip("DuckDB excel extension unavailable in this environment")
    import openpyxl

    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Customers"
    ws1.append(["id", "Full Name"])
    ws1.append([1, "Alice"])
    ws1.append([2, "Bob"])
    ws2 = wb.create_sheet("Region Lookup")
    ws2.append(["code", "region"])
    ws2.append(["NE", "Northeast"])
    p = tmp_path / "workbook.xlsx"
    wb.save(str(p))
    return str(p)


@pytest.fixture
def time_and_datetime_workbook_xlsx(tmp_path, excel_available):
    if not excel_available:
        pytest.skip("DuckDB excel extension unavailable in this environment")
    import datetime

    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["id", "event_time", "event_ts"])
    ws.append([1, datetime.time(14, 30, 0), datetime.datetime(2024, 1, 1, 14, 30, 0)])
    ws.append([2, datetime.time(9, 0, 0), datetime.datetime(2024, 1, 2, 9, 0, 0)])
    p = tmp_path / "time_and_datetime.xlsx"
    wb.save(str(p))
    return str(p)


@pytest.fixture
def headerless_workbook_xlsx(tmp_path, excel_available):
    if not excel_available:
        pytest.skip("DuckDB excel extension unavailable in this environment")
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append([10, 20, 30])
    ws.append([40, 50, 60])
    ws.append([70, 80, 90])
    p = tmp_path / "headerless.xlsx"
    wb.save(str(p))
    return str(p)


@pytest.fixture
def merged_customers(crm_csv, billing_json):
    """
    The worked example, built up to just before AGGREGATE/OUTPUT.

    billing's join field is deliberately left as cust_id (not renamed to
    customer_id) so the merge below has to resolve differently-named keys --
    exactly what a KEY-based MERGE (Form 2) is for.
    """
    crm = rt.read_csv(crm_csv)
    crm = crm.map([
        rt.FieldMapping("id", "customer_id"),
        rt.FieldMapping("full_name", "name"),
    ])
    crm = crm.filter("status = 'active'")

    billing = rt.read_json(billing_json)
    billing = billing.map([
        rt.FieldMapping("amt_due", "balance"),
    ])
    billing = billing.filter("balance > 0")

    return rt.merge([(crm, "customer_id"), (billing, "cust_id")])


#: which column of `one_column_per_type` carries which field type
COLUMN_FOR_TYPE = {
    "bool": "b",
    "int": "i",
    "double": "d",
    "string": "s",
    "date": "dt",
}


@pytest.fixture
def one_column_per_type(tmp_path):
    """A table carrying one column of each of the five field types."""
    p = tmp_path / "all_types.csv"
    p.write_text(
        "b,i,d,s,dt\n"
        "true,1,1.5,hello,2020-01-01\n"
    )
    return rt.read_csv(str(p))


# ---------------------------------------------------------------------------
# normalize_name
# ---------------------------------------------------------------------------

def test_normalize_name_basic():
    assert rt.normalize_name("Field 1") == "field_1"


def test_normalize_name_collapses_internal_spaces():
    assert rt.normalize_name("Field   Num One") == "field_num_one"


def test_normalize_name_trims_leading_and_trailing_whitespace():
    assert rt.normalize_name("  Sheet 1  ") == "sheet_1"


def test_normalize_name_already_normalized_is_unchanged():
    assert rt.normalize_name("full_name") == "full_name"


def test_normalize_name_drops_punctuation_keeping_what_sits_between():
    assert rt.normalize_name("Amount (USD)") == "amount_usd"
    assert rt.normalize_name("Price (GBP) Net") == "price_gbp_net"
    assert rt.normalize_name("(Draft)") == "draft"
    assert rt.normalize_name("Price: USD") == "price_usd"
    assert rt.normalize_name("Total, net") == "total_net"


def test_normalize_name_punctuation_separates_rather_than_vanishing():
    """Without a space of its own, punctuation must not jam words."""
    assert rt.normalize_name("Q1(2024)") == "q1_2024"
    assert rt.normalize_name("a(b)c") == "a_b_c"
    assert rt.normalize_name("Order-Date") == "order_date"
    assert rt.normalize_name("e-mail@work") == "e_mail_work"


def test_normalize_name_drops_a_character_that_leaves_nothing():
    """"$" carries no word of its own, so it leaves no underscore behind."""
    assert rt.normalize_name("Amount ($)") == "amount"
    assert rt.normalize_name("$$$") == "unnamed"


def test_normalize_name_keeps_a_leading_digit():
    """
    A leading digit is legitimate and is neither prefixed nor rejected.
    Grammars are expected to give field and sheet names their own token
    that allows it, rather than the runtime bending the name to fit.
    """
    assert rt.normalize_name("2024 Total") == "2024_total"
    assert rt.normalize_name("1st Quarter") == "1st_quarter"


@pytest.mark.parametrize(
    "raw",
    [
        "Full Name", "Amount ($)", "Order-Date", "Price: USD", "e-mail@work",
        "Q1(2024)", "  Sheet 1  ", "Total, net", "()", "$$$", "%", "  ",
        "50% (net)", "a/b", "col_a", "TOTAL",
    ],
)
def test_normalize_name_always_yields_identifier_characters(raw):
    """
    The result must be usable as an IDENTIFIER token, so it may only ever
    contain letters, digits and underscores, and may never be empty.
    """
    result = rt.normalize_name(raw)
    assert result, f"{raw!r} normalized to an empty name"
    assert re.fullmatch(r"[A-Za-z0-9_]+", result), result
    assert not result.startswith("_") and not result.endswith("_"), result


def test_normalize_name_with_nothing_left_becomes_unnamed():
    """
    A name is the only way to reach a column, so normalizing must never
    produce the empty string.
    """
    assert rt.normalize_name("()") == "unnamed"
    assert rt.normalize_name("   ") == "unnamed"


@pytest.mark.parametrize("raw", [
    "Field 1", "Full Name", "Q1 Results", "Amount (USD)", "Q1(2024)",
    "Order-Date", "Amount ($)", "2024 Total", "a__b", "()", "unnamed",
])
def test_normalize_name_leaves_a_normalized_name_unchanged(raw):
    """
    "In normalized form" means `normalize_name` would return the name
    unchanged. That is what lets the runtime check a program-written name
    with a single call, and what makes a name written into a workbook read
    back under exactly the same name.
    """
    once = rt.normalize_name(raw)
    assert rt.normalize_name(once) == once


def test_read_csv_paren_only_header_is_readable(tmp_path):
    """An empty name would be a zero-length SQL identifier, not a column."""
    csv_path = tmp_path / "parens.csv"
    csv_path.write_text(
        "(),(x)\n"
        "1,2\n"
    )
    table = rt.read_csv(str(csv_path))
    assert list(table.schema) == ["unnamed", "x"]
    assert table._relation.fetchall() == [(1, 2)]


# ---------------------------------------------------------------------------
# can_cast
# ---------------------------------------------------------------------------

@pytest.mark.parametrize("t", sorted(rt.FIELD_TYPES))
def test_can_cast_same_type_is_true(t):
    assert rt.can_cast(t, t) is True


@pytest.mark.parametrize("t", sorted(rt.FIELD_TYPES))
def test_can_cast_to_and_from_string_is_always_true(t):
    assert rt.can_cast(t, "string") is True
    assert rt.can_cast("string", t) is True


@pytest.mark.parametrize(
    "a,b",
    [("bool", "int"), ("int", "bool"),
     ("bool", "double"), ("double", "bool"),
     ("int", "double"), ("double", "int")],
)
def test_can_cast_numeric_family_is_interconvertible(a, b):
    assert rt.can_cast(a, b) is True


@pytest.mark.parametrize("numeric", ["bool", "int", "double"])
def test_can_cast_numeric_never_casts_with_date(numeric):
    assert rt.can_cast(numeric, "date") is False
    assert rt.can_cast("date", numeric) is False


def test_can_cast_rejects_unknown_type_name():
    with pytest.raises(
        rt.PipelineRuntimeError,
        match="'nope' is not one of FIELD_TYPES",
    ):
        rt.can_cast("nope", "int")
    with pytest.raises(
        rt.PipelineRuntimeError,
        match="'nope' is not one of FIELD_TYPES",
    ):
        rt.can_cast("int", "nope")


# ---------------------------------------------------------------------------
# to_sql_type
# ---------------------------------------------------------------------------

def test_to_sql_type_mapping():
    assert rt.to_sql_type("bool") == "BOOLEAN"
    assert rt.to_sql_type("int") == "BIGINT"
    assert rt.to_sql_type("double") == "DOUBLE"
    assert rt.to_sql_type("string") == "VARCHAR"
    assert rt.to_sql_type("date") == "DATE"


def test_to_sql_type_rejects_unknown_type_name():
    with pytest.raises(
        rt.PipelineRuntimeError,
        match="'nope' is not one of FIELD_TYPES",
    ):
        rt.to_sql_type("nope")


# ---------------------------------------------------------------------------
# reading: csv
# ---------------------------------------------------------------------------

def test_read_csv_schema_returns_normalized_names_and_types(crm_csv):
    """
    The schema view and the loaded table agree on every name, so a caller
    never sees a file's raw spelling. The header here is "Full Name".
    """
    assert rt.read_csv_schema(crm_csv) == {
        "id": "int",
        "full_name": "string",
        "status": "string",
        "segment": "string",
    }
    assert list(rt.read_csv_schema(crm_csv)) == list(
        rt.read_csv(crm_csv).schema
    )


def test_read_csv_normalizes_column_names(crm_csv):
    table = rt.read_csv(crm_csv)
    assert table.schema == {
        "id": "int",
        "full_name": "string",
        "status": "string",
        "segment": "string",
    }


def test_read_csv_row_data_is_correct(crm_csv):
    table = rt.read_csv(crm_csv)
    rows = table._relation.fetchall()
    assert len(rows) == 3
    assert rows[0][0] == 1


def test_read_csv_schema_and_read_agree_on_awkward_headers(tmp_path):
    """Normalization is invisible: both views report the same names."""
    csv_path = tmp_path / "awkward.csv"
    csv_path.write_text(
        "  Cust ID  ,Order   Date,TOTAL\n"
        "1,2020-01-01,5\n"
    )
    expected = ["cust_id", "order_date", "total"]
    assert list(rt.read_csv_schema(str(csv_path))) == expected
    assert list(rt.read_csv(str(csv_path)).schema) == expected


def test_read_csv_colliding_names_get_numeric_suffixes(tmp_path):
    """
    Two headers that normalize to one identifier are kept apart with a
    suffix rather than one quietly replacing the other. Both columns, and
    both columns' data, survive.
    """
    csv_path = tmp_path / "collide.csv"
    csv_path.write_text(
        "Full Name,full_name\n"
        "A,B\n"
    )
    expected = ["full_name", "full_name_1"]
    assert list(rt.read_csv_schema(str(csv_path))) == expected

    table = rt.read_csv(str(csv_path))
    assert list(table.schema) == expected
    assert table._relation.fetchall() == [("A", "B")]


def test_read_csv_three_way_collision_numbers_upward(tmp_path):
    csv_path = tmp_path / "three.csv"
    csv_path.write_text(
        "Full Name,FULL NAME,full_name\n"
        "A,B,C\n"
    )
    table = rt.read_csv(str(csv_path))
    assert list(table.schema) == ["full_name", "full_name_1", "full_name_2"]
    assert table._relation.fetchall() == [("A", "B", "C")]


def test_read_csv_suffix_skips_a_name_already_taken(tmp_path):
    """
    A generated suffix must not collide with a real column that already
    carries that name, so it keeps counting up instead.
    """
    csv_path = tmp_path / "taken.csv"
    csv_path.write_text(
        "Full Name,full_name_1,full_name\n"
        "A,B,C\n"
    )
    table = rt.read_csv(str(csv_path))
    assert list(table.schema) == ["full_name", "full_name_1", "full_name_2"]
    assert table._relation.fetchall() == [("A", "B", "C")]


def test_normalized_names_do_not_depend_on_later_columns(tmp_path):
    """Names are assigned left to right, so earlier columns are stable."""
    first = tmp_path / "a.csv"
    first.write_text("Full Name,x\nA,1\n")
    second = tmp_path / "b.csv"
    second.write_text("Full Name,full_name\nA,B\n")
    assert list(rt.read_csv(str(first)).schema)[0] == "full_name"
    assert list(rt.read_csv(str(second)).schema)[0] == "full_name"


def test_read_csv_missing_file_raises(tmp_path):
    with pytest.raises(rt.PipelineRuntimeError, match="could not read CSV"):
        rt.read_csv(str(tmp_path / "nope.csv"))


def test_a_failed_read_leaves_the_connection_usable(tmp_path):
    """
    A statement that fails while executing aborts DuckDB's implicit
    transaction. Left alone, every later statement on the shared connection
    then fails with "Current transaction is aborted" whatever it actually
    asked for. Validation has to report every problem it finds rather than
    stopping at the first, so one failed read must not poison the next.
    """
    good = tmp_path / "good.csv"
    good.write_text(
        "a,b\n"
        "1,2\n"
    )
    missing = str(tmp_path / "nope.csv")

    assert list(rt.read_csv_schema(str(good))) == ["a", "b"]
    for _ in range(3):
        with pytest.raises(
            rt.PipelineRuntimeError,
            match="could not read CSV",
        ):
            rt.read_csv(missing)

    # an unrelated, valid file still reads correctly afterward
    assert list(rt.read_csv_schema(str(good))) == ["a", "b"]
    assert len(rt.read_csv(str(good))._relation.fetchall()) == 1


def test_read_csv_schema_headerless_gets_synthetic_names(headerless_csv):
    assert rt.read_csv_schema(headerless_csv) == {
        "col_a": "int",
        "col_b": "double",
        "col_c": "string",
    }


def test_read_csv_headerless_gets_synthetic_names(headerless_csv):
    table = rt.read_csv(headerless_csv)
    assert table.schema == {
        "col_a": "int",
        "col_b": "double",
        "col_c": "string",
    }
    rows = table._relation.fetchall()
    assert len(rows) == 3
    assert rows[0] == (1, 2.5, "active")


# ---------------------------------------------------------------------------
# reading: json
# ---------------------------------------------------------------------------

def test_read_json_schema(billing_json):
    assert rt.read_json_schema(billing_json) == {
        "cust_id": "int",
        "amt_due": "double",
    }


def test_read_json_row_data_is_correct(billing_json):
    table = rt.read_json(billing_json)
    rows = table._relation.fetchall()
    assert rows == [(1, 50.0), (3, 25.0), (4, 10.0)]


def test_read_json_missing_file_raises(tmp_path):
    with pytest.raises(rt.PipelineRuntimeError, match="could not read JSON"):
        rt.read_json(str(tmp_path / "nope.json"))


def test_read_json_time_and_datetime_widen_to_string(time_and_datetime_json):
    """
    read_json_auto infers TIME/TIMESTAMP for values that look like a
    time-of-day or a combined date+time, both outside FIELD_TYPES now that
    only `date` remains as a temporal type. These must widen to string
    rather than raise PipelineRuntimeError the first time the schema is read.
    """
    schema = rt.read_json_schema(time_and_datetime_json)
    assert schema["event_time"] == "string"
    assert schema["event_ts"] == "string"

    table = rt.read_json(time_and_datetime_json)
    assert table.schema["event_time"] == "string"
    assert table.schema["event_ts"] == "string"
    rows = table._relation.fetchall()
    assert len(rows) == 2


# ---------------------------------------------------------------------------
# reading: xlsx (whole workbook)
# ---------------------------------------------------------------------------

def test_read_excel_workbook_schema_covers_every_sheet(workbook_xlsx):
    schema = rt.read_excel_workbook_schema(workbook_xlsx)
    # sheets are named "Customers" and "Region Lookup" in the workbook
    assert set(schema.keys()) == {"customers", "region_lookup"}
    assert schema["customers"] == {"id": "int", "full_name": "string"}
    assert schema["region_lookup"] == {"code": "string", "region": "string"}


def test_read_excel_workbook_loads_and_normalizes_every_sheet(workbook_xlsx):
    tables = rt.read_excel_workbook(workbook_xlsx)
    assert set(tables.keys()) == {"customers", "region_lookup"}
    assert tables["customers"].schema == {"id": "int", "full_name": "string"}
    assert len(tables["customers"]._relation.fetchall()) == 2

    # the schema view keys the same way, so either can resolve a reference
    assert set(rt.read_excel_workbook_schema(workbook_xlsx)) == set(tables)


def test_read_excel_workbook_colliding_sheet_names_get_suffixes(
    tmp_path, excel_available
):
    """Sheet names collide the same way column names do, and resolve alike."""
    if not excel_available:
        pytest.skip("DuckDB excel extension unavailable in this environment")
    import openpyxl

    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Sales Data"
    ws1.append(["x"])
    ws1.append([1])
    ws2 = wb.create_sheet("sales_data")
    ws2.append(["y"])
    ws2.append([2])
    path = str(tmp_path / "dupe.xlsx")
    wb.save(path)

    tables = rt.read_excel_workbook(path)
    assert list(tables) == ["sales_data", "sales_data_1"]
    assert list(rt.read_excel_workbook_schema(path)) == list(tables)
    assert tables["sales_data"].schema == {"x": "int"}
    assert tables["sales_data_1"].schema == {"y": "int"}


def test_read_excel_workbook_missing_file_raises(tmp_path, excel_available):
    if not excel_available:
        pytest.skip("DuckDB excel extension unavailable in this environment")
    with pytest.raises(
        rt.PipelineRuntimeError,
        match=r"could not read .*nope\.xlsx",
    ):
        rt.read_excel_workbook(str(tmp_path / "nope.xlsx"))


def test_read_excel_workbook_time_and_datetime_widen_to_string(
    time_and_datetime_workbook_xlsx,
):
    """
    read_xlsx reports a time-of-day-formatted cell as TIME and a combined
    date+time cell as TIMESTAMP, both outside FIELD_TYPES now that only
    `date` remains as a temporal type, and unlike CSV there's no
    auto_type_candidates knob to prevent DuckDB from inferring them in the
    first place. These must widen to string rather than raise.
    """
    schema = rt.read_excel_workbook_schema(time_and_datetime_workbook_xlsx)
    assert schema["sheet1"]["event_time"] == "string"
    assert schema["sheet1"]["event_ts"] == "string"

    tables = rt.read_excel_workbook(time_and_datetime_workbook_xlsx)
    table = tables["sheet1"]
    assert table.schema["event_time"] == "string"
    assert table.schema["event_ts"] == "string"
    rows = table._relation.fetchall()
    assert len(rows) == 2


def test_read_excel_workbook_schema_headerless_gets_synthetic_names(
    headerless_workbook_xlsx,
):
    schema = rt.read_excel_workbook_schema(headerless_workbook_xlsx)
    assert schema["sheet1"] == {"col_a": "int", "col_b": "int", "col_c": "int"}


def test_read_excel_workbook_headerless_gets_synthetic_names(headerless_workbook_xlsx):
    tables = rt.read_excel_workbook(headerless_workbook_xlsx)
    table = tables["sheet1"]
    assert table.schema == {"col_a": "int", "col_b": "int", "col_c": "int"}
    rows = table._relation.fetchall()
    assert len(rows) == 3
    assert rows[0] == (10, 20, 30)


# ---------------------------------------------------------------------------
# Table.filter
# ---------------------------------------------------------------------------

def test_filter_keeps_only_matching_rows(crm_csv):
    table = rt.read_csv(crm_csv)
    filtered = table.filter("status = 'active'")
    assert len(filtered._relation.fetchall()) == 2


def test_filter_does_not_mutate_the_original_table(crm_csv):
    table = rt.read_csv(crm_csv)
    before = len(table._relation.fetchall())
    table.filter("status = 'active'")
    after = len(table._relation.fetchall())
    assert before == after == 3


def test_filter_supports_an_explicit_cast(crm_csv):
    table = rt.read_csv(crm_csv)
    filtered = table.filter(f"CAST(id AS {rt.to_sql_type('double')}) > 1.5")
    assert len(filtered._relation.fetchall()) == 2


def test_filter_invalid_sql_raises(crm_csv):
    table = rt.read_csv(crm_csv)
    with pytest.raises(rt.PipelineRuntimeError, match="filter failed"):
        table.filter("this is not valid sql")


# ---------------------------------------------------------------------------
# Table.map
# ---------------------------------------------------------------------------

def test_map_renames_and_leaves_unmentioned_fields_alone(crm_csv):
    table = rt.read_csv(crm_csv)
    mapped = table.map([rt.FieldMapping("id", "customer_id")])
    assert mapped.schema == {
        "customer_id": "int",
        "full_name": "string",
        "status": "string",
        "segment": "string",
    }


def test_map_with_type_changes_the_field_type(crm_csv):
    table = rt.read_csv(crm_csv)
    mapped = table.map([rt.FieldMapping("id", "id", "double")])
    assert mapped.schema["id"] == "double"


def test_map_can_rename_and_cast_in_one_entry(crm_csv):
    table = rt.read_csv(crm_csv)
    mapped = table.map([rt.FieldMapping("id", "customer_id", "double")])
    assert mapped.schema["customer_id"] == "double"
    assert "id" not in mapped.schema


def test_map_unknown_field_raises(crm_csv):
    table = rt.read_csv(crm_csv)
    with pytest.raises(
        rt.PipelineRuntimeError,
        match="map references field 'nonexistent'",
    ):
        table.map([rt.FieldMapping("nonexistent", "x")])


def test_map_illegal_cast_raises(crm_csv):
    table = rt.read_csv(crm_csv)
    with pytest.raises(
        rt.PipelineRuntimeError,
        match="cannot cast field 'id' from 'int' to 'date'",
    ):
        table.map([rt.FieldMapping("id", "id", "date")])


def test_map_unknown_type_alias_raises(crm_csv):
    table = rt.read_csv(crm_csv)
    with pytest.raises(
        rt.PipelineRuntimeError,
        match="not one of this language's five field types",
    ):
        table.map([rt.FieldMapping("id", "id", "nope")])


@pytest.mark.parametrize(
    "from_type,to_type",
    [
        (f, t)
        for f in sorted(rt.FIELD_TYPES)
        for t in sorted(rt.FIELD_TYPES)
        if not rt.can_cast(f, t)
    ],
)
def test_map_raises_whenever_can_cast_is_false(
    from_type, to_type, one_column_per_type
):
    """
    can_cast is the authority on which casts MAP accepts. Whenever it
    reports False, MAP must reject the cast instead of attempting it
    Parametrized off can_cast itself, so a change to the cast
    table cannot leave this check stale.
    """
    col = COLUMN_FOR_TYPE[from_type]
    with pytest.raises(
        rt.PipelineRuntimeError,
        match=f"cannot cast field '{col}' from '{from_type}' to '{to_type}'",
    ):
        one_column_per_type.map([rt.FieldMapping(col, col, to_type)])


def test_map_data_level_cast_failure_raises_from_map_itself(tmp_path):
    """
    can_cast allows string to anything at the type level, so a value that
    does not actually convert can only fail against real data. That is a
    Tier 3 failure and it must be reported by the MAP that caused it,
    as a PipelineRuntimeError rather than a raw DuckDB exception.
    """
    csv_path = tmp_path / "words.csv"
    csv_path.write_text("s\nhello\n")
    table = rt.read_csv(str(csv_path))
    assert table.schema["s"] == "string"
    assert rt.can_cast("string", "int") is True

    with pytest.raises(rt.PipelineRuntimeError, match="map failed"):
        table.map([rt.FieldMapping("s", "s", "int")])


def test_map_cast_failure_in_a_late_row_is_still_caught(tmp_path):
    """
    The check has to scan the whole column, not just the first rows: a
    single bad value at the end must still fail the MAP.
    """
    csv_path = tmp_path / "late.csv"
    rows = "".join(f"{i}\n" for i in range(5000))
    csv_path.write_text("s\n" + rows + "not_a_number\n")
    table = rt.read_csv(str(csv_path))

    with pytest.raises(rt.PipelineRuntimeError, match="not_a_number"):
        table.map([rt.FieldMapping("s", "s", "int")])


def test_map_cast_failure_happens_before_any_later_output(tmp_path):
    """
    A pipeline must run to completion or not at all. Because the
    failing MAP raises immediately, a later OUTPUT never runs and no file
    is left behind for it.
    """
    csv_path = tmp_path / "words.csv"
    csv_path.write_text("s\nhello\n")
    table = rt.read_csv(str(csv_path))
    out = tmp_path / "should_not_exist.csv"

    with pytest.raises(rt.PipelineRuntimeError):
        mapped = table.map([rt.FieldMapping("s", "s", "int")])
        mapped.write_csv(str(out))

    assert not out.exists()


def test_map_without_a_cast_is_unaffected(crm_csv):
    """A plain rename does no cast, so it runs no validation scan."""
    table = rt.read_csv(crm_csv)
    renamed = table.map([rt.FieldMapping("id", "customer_id")])
    assert renamed.schema["customer_id"] == "int"
    assert len(renamed._relation.fetchall()) == 3


# ---------------------------------------------------------------------------
# Table.aggregate
# ---------------------------------------------------------------------------

def test_aggregate_sum_count_min_max_by_group(merged_customers):
    summary = merged_customers.aggregate(
        ["segment"],
        [
            rt.AggregateCall("SUM", "balance", "total_balance"),
            rt.AggregateCall("COUNT", "*", "customer_count"),
            rt.AggregateCall("MAX", "balance", "max_balance"),
        ],
    )
    assert summary.schema == {
        "segment": "string",
        "total_balance": "double",
        "customer_count": "int",
        "max_balance": "double",
    }
    rows = {r[0]: r for r in summary._relation.fetchall()}
    assert rows["enterprise"][1] == 75.0
    assert rows["enterprise"][2] == 2


def test_aggregate_sum_over_int_field_stays_int(tmp_path):
    """
    DuckDB widens SUM over an integral column to HUGEINT, which is outside
    the five field types. The runtime casts it back so the result
    stays describable.
    """
    csv_path = tmp_path / "ints.csv"
    csv_path.write_text(
        "g,n\n"
        "x,10\n"
        "x,20\n"
        "y,5\n"
    )
    table = rt.read_csv(str(csv_path))
    assert table.schema["n"] == "int"

    summary = table.aggregate(["g"], [rt.AggregateCall("SUM", "n", "total")])
    assert summary.schema == {"g": "string", "total": "int"}
    assert dict(summary._relation.fetchall()) == {"x": 30, "y": 5}


def test_aggregate_sum_over_bool_field_stays_int(tmp_path):
    csv_path = tmp_path / "bools.csv"
    csv_path.write_text(
        "g,flag\n"
        "x,true\n"
        "x,false\n"
        "x,true\n"
    )
    table = rt.read_csv(str(csv_path))
    assert table.schema["flag"] == "bool"

    summary = table.aggregate(
        ["g"], [rt.AggregateCall("SUM", "flag", "n_true")]
    )
    assert summary.schema == {"g": "string", "n_true": "int"}
    assert dict(summary._relation.fetchall()) == {"x": 2}


def test_aggregate_avg_is_not_a_supported_function(merged_customers):
    """AVG is deliberately not part of the language."""
    with pytest.raises(rt.PipelineRuntimeError, match="unsupported aggregate"):
        merged_customers.aggregate(
            ["segment"], [rt.AggregateCall("AVG", "balance", "avg_balance")]
        )


def test_aggregate_count_field_counts_non_null(tmp_path):
    """
    COUNT(field) counts non-null values while COUNT(*) counts rows: the
    two forms the language requires of COUNT.
    """
    csv_path = tmp_path / "sparse.csv"
    csv_path.write_text(
        "g,balance\n"
        "x,10\n"
        "x,\n"
        "y,\n"
    )
    table = rt.read_csv(str(csv_path))
    summary = table.aggregate(
        ["g"],
        [
            rt.AggregateCall("COUNT", "balance", "n_values"),
            rt.AggregateCall("COUNT", "*", "n_rows"),
        ],
    )
    assert summary.schema == {
        "g": "string", "n_values": "int", "n_rows": "int",
    }

    rows = {r[0]: (r[1], r[2]) for r in summary._relation.fetchall()}
    assert rows["x"] == (1, 2)  # one non-null balance out of two rows
    assert rows["y"] == (0, 1)  # an all-null group still counts its row


def test_aggregate_star_only_valid_for_count(merged_customers):
    with pytest.raises(
        rt.PipelineRuntimeError,
        match=r"'\*' is only a valid argument for COUNT",
    ):
        merged_customers.aggregate(
            ["segment"], [rt.AggregateCall("SUM", "*", "n")]
        )


def test_aggregate_unknown_function_raises(merged_customers):
    with pytest.raises(
        rt.PipelineRuntimeError,
        match="unsupported aggregate function 'MEDIAN'",
    ):
        merged_customers.aggregate(
            ["segment"], [rt.AggregateCall("MEDIAN", "balance", "n")]
        )


def test_aggregate_missing_group_by_field_raises(merged_customers):
    with pytest.raises(
        rt.PipelineRuntimeError,
        match="aggregate group-by references field 'nonexistent'",
    ):
        merged_customers.aggregate(
            ["nonexistent"], [rt.AggregateCall("COUNT", "*", "n")]
        )


def test_aggregate_missing_argument_field_raises(merged_customers):
    with pytest.raises(
        rt.PipelineRuntimeError,
        match=r"SUM\(nonexistent\) references field 'nonexistent'",
    ):
        merged_customers.aggregate(
            ["segment"], [rt.AggregateCall("SUM", "nonexistent", "n")]
        )


def test_aggregate_multiple_group_by_fields(merged_customers):
    summary = merged_customers.aggregate(
        ["segment", "status"], [rt.AggregateCall("COUNT", "*", "n")]
    )
    assert set(summary.schema) == {"segment", "status", "n"}


# ---------------------------------------------------------------------------
# Normalization reaches columns and sheets, and stops there
# ---------------------------------------------------------------------------

def test_file_paths_are_never_normalized(tmp_path):
    """
    Normalizing a path would send a read looking for a file that isn't
    there. Only column names and sheet names are normalized.
    """
    awkward = tmp_path / "My Data (2024).csv"
    awkward.write_text(
        "Full Name,x\n"
        "A,1\n"
    )

    # the path is honored verbatim, while the header inside it normalizes
    assert list(rt.read_csv(str(awkward)).schema) == ["full_name", "x"]
    assert list(rt.read_csv_schema(str(awkward))) == ["full_name", "x"]


def test_output_paths_are_never_normalized(tmp_path):
    table = rt.read_csv(str(_awkward_source(tmp_path)))

    out_csv = tmp_path / "Q1 Report (Final).csv"
    table.write_csv(str(out_csv))
    assert out_csv.exists()

    out_json = tmp_path / "Q1 Report (Final).json"
    table.write_json(str(out_json))
    assert out_json.exists()


def test_workbook_path_is_not_normalized(tmp_path, excel_available):
    if not excel_available:
        pytest.skip("DuckDB excel extension unavailable in this environment")
    table = rt.read_csv(str(_awkward_source(tmp_path)))

    out = tmp_path / "My Workbook (v2).xlsx"
    rt.write_excel_workbook({"q1_results": table}, str(out))

    assert out.exists()  # path verbatim
    assert list(rt.read_excel_workbook(str(out))) == ["q1_results"]


def _awkward_source(tmp_path):
    p = tmp_path / "Source File (raw).csv"
    p.write_text(
        "Full Name,x\n"
        "A,1\n"
    )
    return p


# ---------------------------------------------------------------------------
# Program-supplied names are never normalized: a name a program creates must
# already be in normalized form, and the runtime rejects one that isn't
# rather than quietly fixing it or using it verbatim
# ---------------------------------------------------------------------------

def test_map_uses_a_normalized_target_name_exactly(crm_csv):
    table = rt.read_csv(crm_csv)
    mapped = table.map([rt.FieldMapping("id", "customer_id")])
    assert list(mapped.schema)[0] == "customer_id"

    # a leading digit and a doubled underscore are both normalized form
    mapped = table.map([rt.FieldMapping("id", "2024__total")])
    assert list(mapped.schema)[0] == "2024__total"
    assert list(mapped.select(["2024__total"]).schema) == ["2024__total"]


@pytest.mark.parametrize("target", [
    "Customer ID", "customer id", "CUSTOMER_ID", "Amount (USD)",
    "_customer_id", "customer_id_", "",
])
def test_map_rejects_a_target_name_not_in_normalized_form(crm_csv, target):
    """
    The grammar is meant to make this unreachable. If it doesn't, the
    runtime must not normalize on the program's behalf, since that would
    hide the grammar bug, and must not create a column under a name the
    grammar could never refer to.
    """
    table = rt.read_csv(crm_csv)
    with pytest.raises(
        rt.PipelineRuntimeError, match="grammar should have rejected"
    ):
        table.map([rt.FieldMapping("id", target)])
    assert list(table.schema) == ["id", "full_name", "status", "segment"]


def test_map_does_not_normalize_the_name_it_refers_to(crm_csv):
    """
    A reference must already be the normalized name. The CSV header is
    "Full Name", so the field is full_name and only that spelling finds it.
    Working that out is the caller's job, which is why the rule matters.
    """
    table = rt.read_csv(crm_csv)

    mapped = table.map([rt.FieldMapping("full_name", "name")])
    assert "name" in mapped.schema

    for spelling in ["Full Name", "full name", "FULL_NAME"]:
        with pytest.raises(rt.PipelineRuntimeError, match="does not exist"):
            table.map([rt.FieldMapping(spelling, "name")])


def test_aggregate_rejects_an_output_name_not_in_normalized_form(
    merged_customers,
):
    """The result name is created, so it must already be normalized."""
    with pytest.raises(
        rt.PipelineRuntimeError, match="grammar should have rejected"
    ):
        merged_customers.aggregate(
            ["segment"], [rt.AggregateCall("SUM", "balance", "Total Balance")]
        )
    summary = merged_customers.aggregate(
        ["segment"], [rt.AggregateCall("SUM", "balance", "total_balance")]
    )
    assert list(summary.schema) == ["segment", "total_balance"]


def test_aggregate_does_not_normalize_the_fields_it_refers_to(
    merged_customers,
):
    with pytest.raises(rt.PipelineRuntimeError, match="does not exist"):
        merged_customers.aggregate(
            ["Segment"], [rt.AggregateCall("SUM", "balance", "total")]
        )
    with pytest.raises(rt.PipelineRuntimeError, match="does not exist"):
        merged_customers.aggregate(
            ["segment"], [rt.AggregateCall("SUM", "Balance", "total")]
        )


def test_select_does_not_normalize_the_names_it_refers_to(crm_csv):
    table = rt.read_csv(crm_csv)
    assert list(table.select(["full_name", "id"]).schema) == [
        "full_name",
        "id",
    ]
    with pytest.raises(rt.PipelineRuntimeError, match="does not exist"):
        table.select(["Full Name"])


def test_merge_does_not_normalize_its_key_names(crm_csv, billing_json):
    crm = rt.read_csv(crm_csv).map([rt.FieldMapping("id", "customer_id")])
    billing = rt.read_json(billing_json).map(
        [rt.FieldMapping("cust_id", "customer_id")]
    )
    merged = rt.merge([(crm, "customer_id"), (billing, "customer_id")])
    assert len(merged._relation.fetchall()) == 2

    with pytest.raises(rt.PipelineRuntimeError, match="merge key"):
        rt.merge([(crm, "Customer ID"), (billing, "customer_id")])


def test_write_excel_workbook_rejects_a_sheet_name_not_in_normalized_form(
    crm_csv, tmp_path, excel_available
):
    """
    A workbook this program writes must be referenceable by the same
    program under exactly the names it wrote, so a sheet name is held to
    the same rule as a MAP target: already normalized, or rejected before
    anything is written.
    """
    if not excel_available:
        pytest.skip("DuckDB excel extension unavailable in this environment")
    table = rt.read_csv(crm_csv)
    out = tmp_path / "report.xlsx"
    with pytest.raises(
        rt.PipelineRuntimeError, match="grammar should have rejected"
    ):
        rt.write_excel_workbook({"Q1 Results": table}, str(out))
    assert not out.exists()

    rt.write_excel_workbook({"q1_results": table}, str(out))
    assert list(rt.read_excel_workbook(str(out))) == ["q1_results"]


def test_a_real_typo_still_raises(crm_csv):
    """Normalizing references must not paper over a genuine mistake."""
    table = rt.read_csv(crm_csv)
    with pytest.raises(rt.PipelineRuntimeError, match="does not exist"):
        table.select(["nope"])


# ---------------------------------------------------------------------------
# Table.select
# ---------------------------------------------------------------------------

def test_select_keeps_only_named_columns_in_given_order(crm_csv):
    table = rt.read_csv(crm_csv)
    narrowed = table.select(["segment", "id"])
    assert list(narrowed.schema.keys()) == ["segment", "id"]


def test_select_row_data_is_correct(crm_csv):
    table = rt.read_csv(crm_csv)
    narrowed = table.select(["id", "status"])
    rows = narrowed._relation.fetchall()
    assert rows == [(1, "active"), (2, "inactive"), (3, "active")]


def test_select_does_not_mutate_the_original_table(crm_csv):
    table = rt.read_csv(crm_csv)
    table.select(["id"])
    assert list(table.schema.keys()) == ["id", "full_name", "status", "segment"]


def test_select_unknown_column_raises(crm_csv):
    table = rt.read_csv(crm_csv)
    with pytest.raises(
        rt.PipelineRuntimeError,
        match="select references field 'nonexistent'",
    ):
        table.select(["nonexistent"])


def test_select_then_write_csv_narrows_the_file(crm_csv, tmp_path):
    table = rt.read_csv(crm_csv)
    out = str(tmp_path / "out.csv")
    table.select(["segment", "id"]).write_csv(out)
    header = open(out).readline().strip()
    assert header == "segment,id"


# ---------------------------------------------------------------------------
# Table.write_csv / write_json (single-table)
# ---------------------------------------------------------------------------

def test_write_csv_all_columns_round_trips(crm_csv, tmp_path):
    table = rt.read_csv(crm_csv)
    out = str(tmp_path / "out.csv")
    table.write_csv(out)
    back = rt.read_csv(out)
    assert back.schema == table.schema
    assert len(back._relation.fetchall()) == 3


def test_write_json_all_columns_round_trips(billing_json, tmp_path):
    table = rt.read_json(billing_json)
    out = str(tmp_path / "out.json")
    table.write_json(out)
    back = rt.read_json(out)
    assert back.schema == table.schema
    assert len(back._relation.fetchall()) == 3


# ---------------------------------------------------------------------------
# merge
# ---------------------------------------------------------------------------

def test_merge_basic_two_table_join(crm_csv, billing_json):
    crm = rt.read_csv(crm_csv).map([rt.FieldMapping("id", "customer_id")])
    billing = rt.read_json(billing_json).map(
        [rt.FieldMapping("cust_id", "customer_id")]
    )
    merged = rt.merge([(crm, "customer_id"), (billing, "customer_id")])
    rows = merged._relation.fetchall()
    # only customers 1 and 3 appear in both crm and billing
    assert len(rows) == 2


def test_merge_key_output_name_comes_from_first_table(crm_csv, billing_json):
    crm = rt.read_csv(crm_csv).map([rt.FieldMapping("id", "the_key")])
    billing = rt.read_json(billing_json).map(
        [rt.FieldMapping("cust_id", "totally_different_name")]
    )
    merged = rt.merge(
        [(crm, "the_key"), (billing, "totally_different_name")]
    )
    assert "the_key" in merged.schema
    assert "totally_different_name" not in merged.schema


def test_merge_three_tables(crm_csv, billing_json, tmp_path):
    regions_csv = tmp_path / "regions.csv"
    regions_csv.write_text("cid,region\n1,Northeast\n3,South\n")

    crm = rt.read_csv(crm_csv).map([rt.FieldMapping("id", "customer_id")])
    billing = rt.read_json(billing_json).map(
        [rt.FieldMapping("cust_id", "customer_id")]
    )
    regions = rt.read_csv(str(regions_csv)).map(
        [rt.FieldMapping("cid", "customer_id")]
    )
    merged = rt.merge(
        [(crm, "customer_id"), (billing, "customer_id"), (regions, "customer_id")]
    )
    assert "region" in merged.schema
    assert len(merged._relation.fetchall()) == 2


def test_merge_requires_at_least_two_bindings(crm_csv):
    crm = rt.read_csv(crm_csv)
    with pytest.raises(
        rt.PipelineRuntimeError,
        match="merge requires at least two bindings",
    ):
        rt.merge([(crm, "id")])


def test_merge_missing_key_field_raises(crm_csv, billing_json):
    crm = rt.read_csv(crm_csv)
    billing = rt.read_json(billing_json)
    with pytest.raises(
        rt.PipelineRuntimeError,
        match="merge key 'nonexistent' does not exist",
    ):
        rt.merge([(crm, "id"), (billing, "nonexistent")])


def test_merge_non_key_column_collision_raises(crm_csv, billing_json):
    # both tables have a field called "status" that isn't the join key
    crm = rt.read_csv(crm_csv).map([rt.FieldMapping("id", "customer_id")])
    billing = rt.read_json(billing_json).map([
        rt.FieldMapping("cust_id", "customer_id"),
        rt.FieldMapping("amt_due", "status"),  # deliberately collides
    ])
    with pytest.raises(
        rt.PipelineRuntimeError,
        match="non-key column collision",
    ):
        rt.merge([(crm, "customer_id"), (billing, "customer_id")])


# ---------------------------------------------------------------------------
# write_excel_workbook (multi-table)
# ---------------------------------------------------------------------------

def test_write_excel_workbook_round_trips(crm_csv, billing_json, tmp_path, excel_available):
    if not excel_available:
        pytest.skip("DuckDB excel extension unavailable in this environment")
    crm = rt.read_csv(crm_csv)
    billing = rt.read_json(billing_json)
    out = str(tmp_path / "report.xlsx")
    rt.write_excel_workbook({"customers": crm, "billing": billing}, out)

    # already-normalized sheet names read back unchanged
    back = rt.read_excel_workbook(out)
    assert set(back.keys()) == {"customers", "billing"}
    assert len(back["customers"]._relation.fetchall()) == 3
    assert len(back["billing"]._relation.fetchall()) == 3


def test_write_excel_workbook_requires_at_least_one_sheet(excel_available):
    if not excel_available:
        pytest.skip("DuckDB excel extension unavailable in this environment")
    with pytest.raises(
        rt.PipelineRuntimeError,
        match="requires at least one sheet",
    ):
        rt.write_excel_workbook({}, "/tmp/should_not_be_created.xlsx")


def test_write_excel_workbook_is_atomic_on_failure(
    crm_csv, tmp_path, excel_available
):
    """
    If a later sheet fails to write, the target path must be left exactly as
    it was before the call: no partially-written file (the atomicity
    guarantee for the multi-sheet writer).
    """
    if not excel_available:
        pytest.skip("DuckDB excel extension unavailable in this environment")

    good = rt.read_csv(crm_csv)

    # A table whose underlying file disappears before the write actually
    # happens. Forces a real failure partway through the sheet sequence,
    # without needing to mock anything.
    doomed_path = tmp_path / "doomed.csv"
    doomed_path.write_text("x\n1\n")
    doomed = rt.read_csv(str(doomed_path))
    os.remove(doomed_path)

    out = tmp_path / "atomic.xlsx"
    out.write_text("pre-existing content that must survive")

    with pytest.raises(

        rt.PipelineRuntimeError,

        match="could not write workbook",

    ):
        rt.write_excel_workbook({"good": good, "doomed": doomed}, str(out))

    # original file untouched
    assert out.read_text() == "pre-existing content that must survive"
    # no stray temp files left in the directory
    leftovers = [f for f in os.listdir(tmp_path) if f not in {
        "crm.csv", "doomed.csv", "atomic.xlsx"
    }]
    assert leftovers == []


# ---------------------------------------------------------------------------
# End-to-end: the worked example, in full
# ---------------------------------------------------------------------------

def test_full_worked_example_end_to_end(crm_csv, billing_json, tmp_path):
    crm = rt.read_csv(crm_csv)
    crm = crm.map([
        rt.FieldMapping("id", "customer_id"),
        rt.FieldMapping("full_name", "name"),
    ])
    crm = crm.filter("status = 'active'")

    billing = rt.read_json(billing_json)
    billing = billing.map([
        rt.FieldMapping("amt_due", "balance"),
    ])
    billing = billing.filter("balance > 0")

    # billing's join field stays cust_id, resolved via a KEY (Form 2),
    # matching the worked example rather than renaming both sides.
    customers = rt.merge([(crm, "customer_id"), (billing, "cust_id")])
    unified_out = str(tmp_path / "unified_customers.json")
    customers.write_json(unified_out)

    summary = customers.aggregate(
        ["segment"],
        [
            rt.AggregateCall("SUM", "balance", "total_balance"),
            rt.AggregateCall("COUNT", "*", "customer_count"),
            rt.AggregateCall("MAX", "balance", "max_balance"),
        ],
    )
    summary_out = str(tmp_path / "segment_summary.csv")
    summary_slim = summary.select(["segment", "total_balance", "customer_count"])
    summary_slim.write_csv(summary_out)

    assert os.path.exists(unified_out)
    header = open(summary_out).readline().strip()
    assert header == "segment,total_balance,customer_count"

    back = rt.read_json(unified_out)
    assert len(back._relation.fetchall()) == 2
