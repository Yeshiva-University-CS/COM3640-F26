"""Compare a successful test output against the reference expectations."""
import argparse
import csv
import datetime
import json
import math
from pathlib import Path


def equal(actual, expected, csv_mode=False):
    if expected is None:
        return actual == '' if csv_mode else actual is None
    if isinstance(expected, bool):
        return str(actual).lower() == str(expected).lower() if csv_mode else actual is expected
    if isinstance(expected, (int, float)):
        if isinstance(actual, bool) or (not csv_mode and not isinstance(actual, (int, float))):
            return False
        try:
            return math.isclose(float(actual), expected, rel_tol=1e-9, abs_tol=1e-9)
        except (TypeError, ValueError):
            return False
    if isinstance(actual, (datetime.date, datetime.datetime)):
        actual = actual.isoformat()
    return actual == expected


def compare(table, expected, csv_mode=False):
    columns, rows = table
    if columns != expected['columns']:
        raise ValueError(f"columns: expected {expected['columns']!r}; got {columns!r}")
    remaining = list(rows)
    for wanted in expected['rows']:
        for index, actual in enumerate(remaining):
            if len(actual) == len(wanted) and all(equal(a, e, csv_mode) for a, e in zip(actual, wanted)):
                remaining.pop(index)
                break
        else:
            raise ValueError(f'missing or incorrect row: {wanted!r}')
    if remaining:
        raise ValueError(f'unexpected extra rows: {remaining!r}')


def check(case, path):
    expected = json.loads((Path(__file__).resolve().parent / 'expected.json').read_text(encoding='utf-8'))[case]
    if 'sheets' in expected:
        import openpyxl
        book = openpyxl.load_workbook(path, data_only=True, read_only=True)
        try:
            if set(book.sheetnames) != set(expected['sheets']):
                raise ValueError(f"expected sheets {list(expected['sheets'])}; got {book.sheetnames}")
            for name, table in expected['sheets'].items():
                rows = list(book[name].values)
                compare((list(rows[0]), [list(row) for row in rows[1:]]), table)
        finally:
            book.close()
    elif path.suffix.lower() == '.json':
        raw = path.read_text(encoding='utf-8-sig')
        try:
            rows = json.loads(raw)
        except json.JSONDecodeError:
            rows = [json.loads(line) for line in raw.splitlines() if line.strip()]
        if not isinstance(rows, list) or not all(isinstance(row, dict) for row in rows):
            raise ValueError('JSON must contain records')
        columns = expected['columns']
        if any(set(row) != set(columns) for row in rows):
            raise ValueError('incorrect JSON field names')
        compare((columns, [[row[col] for col in columns] for row in rows]), expected)
    else:
        with path.open(newline='', encoding='utf-8-sig') as f:
            rows = list(csv.reader(f))
        if not rows:
            raise ValueError('empty file without column headers')
        compare((rows[0], rows[1:]), expected, csv_mode=True)


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('case', help='artifact ID, e.g. 01-original or 03-workbook')
    parser.add_argument('output', type=Path)
    args = parser.parse_args()
    try:
        # Enforce the specifically requested output formats.
        required = '.json' if args.case == '02-json' else '.xlsx' if args.case == '03-workbook' else '.csv'
        if args.output.suffix.lower() != required:
            raise ValueError(f'{args.case} requires {required} output')
        check(args.case, args.output)
    except (ValueError, KeyError, OSError, IndexError) as error:
        parser.exit(1, f'FAIL {args.case}: {error}\n')
    print(f'PASS {args.case}')
