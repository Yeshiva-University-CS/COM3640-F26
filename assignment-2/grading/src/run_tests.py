"""Run the ten PipelineQL flows in isolated working directories."""
from __future__ import annotations

import argparse
import contextlib
import csv
import datetime
from fractions import Fraction
import io
import json
from pathlib import Path
import re
import shutil
import subprocess
import sys
import tempfile
import traceback
import uuid

from check_output import compare

HERE = Path(__file__).resolve().parent.parent
ROOT = HERE.parent


def exception_info(error):
    result = {'type': type(error).__name__, 'message': str(error)}
    if getattr(error, '__cause__', None) is not None:
        result['cause'] = type(error.__cause__).__name__
    if hasattr(error, 'errors'):
        result['errors'] = [exception_info(item) for item in error.errors]
    return result


def worker(workspace, program, result_path):
    """Observe the real driver, without altering its error handling."""
    sys.path.insert(0, str(workspace))
    result = {'stage': 'setup', 'parsed': False, 'validated': False,
              'executed': False, 'errors': []}
    stdout, stderr = io.StringIO(), io.StringIO()
    try:
        with contextlib.redirect_stdout(stdout), contextlib.redirect_stderr(stderr):
            import main as engine
            original_load = engine._load_interpreter
            original_build = engine._build_parser

            def load(required):
                module = original_load(required)
                original_validate, original_execute = module.validate, module.execute

                def validate(tree):
                    result['stage'] = 'validation'
                    try:
                        original_validate(tree)
                    except Exception as error:
                        result['validation_error'] = exception_info(error)
                        raise
                    result['validated'] = True

                def execute(tree):
                    result['stage'] = 'execution'
                    try:
                        original_execute(tree)
                    except Exception as error:
                        result['execution_error'] = exception_info(error)
                        raise
                    result['executed'] = True

                module.validate, module.execute = validate, execute
                return module

            def build(start, parser):
                grammar = original_build(start, parser)

                class ObservedParser:
                    def parse(self, source):
                        result['stage'] = 'parse'
                        try:
                            tree = grammar.parse(source)
                        except Exception as error:
                            result['parse_error'] = exception_info(error)
                            result['lark_parse_error'] = isinstance(error, engine.lark.exceptions.LarkError)
                            raise
                        result['parsed'] = True
                        return tree

                return ObservedParser()

            engine._load_interpreter, engine._build_parser = load, build
            result['exit_code'] = engine.main([str(program)])
    except BaseException as error:
        result['uncaught'] = exception_info(error)
        result['traceback'] = traceback.format_exc()
        result['exit_code'] = 2
    result['stdout'], result['stderr'] = stdout.getvalue(), stderr.getvalue()
    result_path.write_text(json.dumps(result, indent=2), encoding='utf-8')


def read_table(path, expected, sheet=None):
    if path.suffix == '.xlsx':
        import openpyxl
        book = openpyxl.load_workbook(path, data_only=True, read_only=True)
        try:
            rows = list(book[sheet].values)
            return (list(rows[0]), [list(row) for row in rows[1:]]), False
        finally:
            book.close()
    if path.suffix == '.json':
        raw = path.read_text(encoding='utf-8-sig')
        try:
            records = json.loads(raw)
        except json.JSONDecodeError:
            records = [json.loads(line) for line in raw.splitlines() if line.strip()]
        if not isinstance(records, list) or not all(isinstance(row, dict) for row in records):
            raise ValueError('JSON must contain records')
        columns = expected['columns']
        if any(set(row) != set(columns) for row in records):
            raise ValueError('incorrect JSON fields')
        return (columns, [[row[col] for col in columns] for row in records]), False
    with path.open(newline='', encoding='utf-8-sig') as stream:
        rows = list(csv.reader(stream))
    if not rows:
        raise ValueError('missing header')
    return (rows[0], rows[1:]), True


def score_flow(flow, observed, output, expected):
    checks = []
    validation = observed.get('validation_error', {})
    errors = validation.get('errors', []) if validation.get('type') == 'ValidationErrors' else []
    used_errors = set()
    no_output = not any(output.rglob('*'))
    setup = observed.get('stage') == 'setup' or observed.get('timeout', False)

    def grade(check):
        kind = check['kind']
        if setup:
            return 'BLOCKED', observed.get('problem', 'setup failure; see log')
        if kind in {'table', 'sheet', 'projection'}:
            path = output / check['file']
            if not path.is_file():
                if observed.get('exit_code') != 0:
                    return 'BLOCKED', 'output not reached after earlier failure; see log'
                return 'FAIL', 'required output file missing'
            reference = expected[check['id']]
            if kind == 'sheet':
                reference = reference['sheets'][check['sheet']]
            try:
                actual, csv_mode = read_table(path, reference, check.get('sheet'))
                if kind == 'sheet':
                    import openpyxl
                    book = openpyxl.load_workbook(path, read_only=True)
                    try:
                        if set(book.sheetnames) != set(expected[check['id']]['sheets']):
                            raise ValueError('workbook must contain exactly the requested sheets')
                    finally:
                        book.close()
                if kind == 'projection':
                    columns, rows = actual
                    wanted = check['columns']
                    if len(columns) != len(set(columns)):
                        raise ValueError('duplicate column names')
                    indices = [columns.index(name) for name in wanted]
                    if indices != sorted(indices):
                        raise ValueError('incorrect column order')
                    actual = (wanted, [[row[index] for index in indices] for row in rows])
                    ref_indices = [reference['columns'].index(name) for name in wanted]
                    reference = {'columns': wanted, 'rows': [[row[i] for i in ref_indices] for row in reference['rows']]}
                compare(actual, reference, csv_mode)
                return 'PASS', 'output matches'
            except Exception as error:
                return 'FAIL', str(error)
        if kind == 'parse_rejection':
            return ('PASS', 'grammar rejected program') if observed.get('lark_parse_error') else ('FAIL', 'no grammar rejection')
        if kind == 'driver_error':
            ok = observed.get('exit_code') == 1 and 'GrammarError:' in observed.get('stderr', '') and no_output
            return ('PASS', 'driver reports GrammarError; no output') if ok else ('FAIL', 'incorrect driver result or output produced')
        if kind == 'validation_issue':
            if not observed.get('parsed'):
                return 'BLOCKED', 'validation not reached'
            for index, error in enumerate(errors):
                if index in used_errors or error['type'] != check['exception']:
                    continue
                if re.search(r'(?<!\w)' + re.escape(check['identifier']) + r'(?!\w)', error['message'], re.IGNORECASE):
                    used_errors.add(index)
                    return 'PASS', 'distinct diagnostic identifies ' + check['identifier']
            if errors:
                if not any(error['type'] == check['exception'] for error in errors):
                    return 'FAIL', 'validation errors have the wrong tier'
                return 'BLOCKED', 'no distinct matching diagnostic; early stop or unidentified issue requires separate test'
            return 'FAIL', 'expected validation error not reported'
        if kind == 'validation_exact':
            ok = (len(errors) == check['count'] and
                  all(error['type'] == check['exception'] for error in errors) and
                  len(used_errors) == check['count'] and
                  observed.get('exit_code') == check.get('exit_code', 1) and no_output)
            return ('PASS', 'all distinct errors collected') if ok else ('FAIL', 'error collection incomplete or duplicated (single collection deduction)')
        if kind == 'validation_rejection':
            ok = (observed.get('parsed') and errors and all(error['type'] == check['exception'] for error in errors)
                  and observed.get('exit_code') == 1 and no_output)
            return ('PASS', 'correct validation tier; no output') if ok else ('FAIL', 'wrong rejection tier/exit or output produced')
        if kind == 'validation_success':
            return ('PASS', 'validation succeeded') if observed.get('validated') else ('FAIL', 'validation did not succeed')
        if kind == 'execution_error':
            error = observed.get('execution_error', {})
            ok = (observed.get('validated') and error.get('type') == check['exception'] and
                  error.get('cause') == check['cause'] and observed.get('exit_code') == 1 and no_output)
            return ('PASS', 'wrapped runtime failure; no output') if ok else ('FAIL', 'incorrect runtime exception, cause, exit, or output')
        return 'FAIL', 'unknown checkpoint kind: ' + kind

    for check in flow['checks']:
        status, detail = grade(check)
        label = check.get('file', check['kind'])
        if 'columns' in check:
            label += ':' + ','.join(check['columns'])
        if 'identifier' in check:
            label += ':' + check['identifier']
        if 'sheet' in check:
            label += ':' + check['sheet']
        checks.append({'check': label, 'points': check['points'], 'status': status, 'detail': detail})
    # Reject extra/reordered summary columns without destroying unrelated projections.
    if flow['program'] == 'flow05.ql' and all(c['status'] == 'PASS' for c in checks):
        try:
            with (output / '05-region.csv').open(newline='', encoding='utf-8-sig') as stream:
                columns = next(csv.reader(stream))
            if columns != expected['05-region']['columns']:
                checks[0].update(status='FAIL', detail='incorrect full summary column list/order')
        except Exception as error:
            checks[0].update(status='FAIL', detail=str(error))
    return checks


def write_score_matrix(report, destination):
    rows = []
    for flow in report['flows']:
        maximum = sum((Fraction(c['points']) for c in flow['checks']), Fraction())
        score = sum((Fraction(c['points']) for c in flow['checks'] if c['status'] == 'PASS'), Fraction())
        rows.append({'flow': Path(flow['program']).stem, 'value': float(maximum), 'score': float(score)})
    rows.append({'flow': 'Total', 'value': report['maximum_score'], 'score': report['earned']})
    report['score_matrix'] = rows
    with (destination / 'scores.csv').open('w', newline='', encoding='utf-8') as stream:
        writer = csv.writer(stream)
        writer.writerow(['flow', 'value', 'score'])
        writer.writerows([row['flow'], f"{row['value']:g}", f"{row['score']:.2f}"] for row in rows)
    matrix = '\n'.join([f"{'flow':<10} {'value':>7} {'score':>7}",
                        *[f"{row['flow']:<10} {row['value']:>7g} {row['score']:>7.2f}" for row in rows]])
    return matrix


def run(args):
    manifest = json.loads((HERE / 'src' / 'automation.json').read_text(encoding='utf-8'))
    expected = json.loads((HERE / 'src' / 'expected.json').read_text(encoding='utf-8'))
    flows = manifest['flows']
    if args.flow:
        flows = [f for f in flows if int(f['program'][4:6]) in args.flow]
    stamp = datetime.datetime.now().strftime('%Y%m%d-%H%M%S') + '-' + uuid.uuid4().hex[:6]
    destination = args.results.resolve() / stamp
    destination.mkdir(parents=True)
    report = {'maximum_score': sum(float(sum(Fraction(c['points']) for c in f['checks'])) for f in flows), 'flows': []}
    for flow in flows:
        name = Path(flow['program']).stem
        saved = destination / name
        saved.mkdir()
        source = args.implementation.resolve()
        program = source / 'grading' / 'ql' / flow['program']
        if not program.is_file():
            checks = [{'check': 'program', 'points': '2', 'status': 'FAIL', 'detail': f'missing {program}'}]
            observed = {'problem': 'program missing'}
        else:
            with tempfile.TemporaryDirectory(prefix='pipelineql-') as temporary:
                workspace = Path(temporary)
                (workspace / 'grading' / 'output').mkdir(parents=True)
                shutil.copytree(HERE / 'data', workspace / 'grading' / 'data')
                shutil.copytree(source / 'grading' / 'ql', workspace / 'grading' / 'ql')
                for filename in ['main.py', 'pipeline_runtime.py', 'pipeline_errors.py']:
                    shutil.copy2(ROOT / filename, workspace / filename)
                missing = []
                for filename in ['pipelineql.lark', 'interpreter.py']:
                    if (source / filename).is_file():
                        shutil.copy2(source / filename, workspace / filename)
                    else:
                        missing.append(filename)
                result_path = workspace / 'observation.json'
                if missing:
                    observed = {'stage': 'setup', 'problem': 'missing ' + ', '.join(missing)}
                else:
                    try:
                        process = subprocess.run(
                            [sys.executable, str(Path(__file__).resolve()), '--worker', str(workspace),
                             str(Path('grading/ql') / flow['program']), str(result_path)],
                            cwd=workspace, capture_output=True, text=True, encoding='utf-8', errors='replace', timeout=args.timeout,
                        )
                        observed = json.loads(result_path.read_text(encoding='utf-8')) if result_path.is_file() else {
                            'stage': 'setup', 'problem': 'worker failed', 'stdout': process.stdout, 'stderr': process.stderr}
                    except subprocess.TimeoutExpired as error:
                        observed = {'stage': 'setup', 'timeout': True, 'problem': f'timeout after {args.timeout:g}s'}
                output = workspace / 'grading' / 'output'
                checks = score_flow(flow, observed, output, expected)
                shutil.copytree(output, saved / 'output')
        (saved / 'observation.json').write_text(json.dumps(observed, indent=2), encoding='utf-8')
        (saved / 'log.txt').write_text(observed.get('stdout', '') + '\n' + observed.get('stderr', '') + '\n' + observed.get('traceback', ''), encoding='utf-8')
        report['flows'].append({'program': flow['program'], 'checks': checks})
        earned = sum(Fraction(c['points']) for c in checks if c['status'] == 'PASS')
        blocked = sum(Fraction(c['points']) for c in checks if c['status'] == 'BLOCKED')
        print(f'{name}: {float(earned):.2f}/2' + (f' ({float(blocked):.2f} blocked)' if blocked else ''))
        for check in checks:
            print(f"  {check['status']:7} {check['check']}: {check['detail']}")
    all_checks = [c for f in report['flows'] for c in f['checks']]
    earned = sum((Fraction(c['points']) for c in all_checks if c['status'] == 'PASS'), Fraction())
    blocked = sum((Fraction(c['points']) for c in all_checks if c['status'] == 'BLOCKED'), Fraction())
    report.update(earned=float(earned), blocked=float(blocked), complete=not blocked,
                  possible_score_range=[float(earned), float(earned + blocked)])
    summary = f"Confirmed: {float(earned):.2f}/{report['maximum_score']:g}; blocked: {float(blocked):.2f}"
    if blocked:
        summary += f'; unresolved score range {float(earned):.2f}-{float(earned + blocked):.2f}'
    matrix = write_score_matrix(report, destination)
    print('\n' + matrix)
    print(summary)
    print('Results:', destination)
    (destination / 'report.json').write_text(json.dumps(report, indent=2), encoding='utf-8')
    (destination / 'summary.txt').write_text(matrix + '\n\n' + summary + '\n', encoding='utf-8')
    return 0 if all(c['status'] == 'PASS' for c in all_checks) else 2 if blocked else 1


def main():
    if len(sys.argv) > 1 and sys.argv[1] == '--worker':
        worker(Path(sys.argv[2]), Path(sys.argv[3]), Path(sys.argv[4]))
        return 0
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--implementation', type=Path, default=ROOT, help='directory containing interpreter.py, grammar, and grading/ql/')
    parser.add_argument('--results', type=Path, default=HERE / 'results')
    parser.add_argument('--timeout', type=float, default=60)
    parser.add_argument('--flow', type=int, nargs='+', choices=range(1, 11))
    args = parser.parse_args()
    if args.timeout <= 0:
        parser.error('--timeout must be positive')
    return run(args)


if __name__ == '__main__':
    sys.exit(main())
